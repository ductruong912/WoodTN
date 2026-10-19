"""Evaluate a PyTorch checkpoint and write a classification report and matrix."""

import argparse
import csv
import json
from pathlib import Path

import torch
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from sklearn.metrics import classification_report, confusion_matrix
from torch.utils.data import DataLoader
from torchvision.datasets import ImageFolder
from tqdm import tqdm

from tools.model_utils import (
    ARCHITECTURES,
    DEFAULT_CLASSES,
    evaluation_transform,
    load_model,
    read_classes,
)


def write_excel_report(path, report, matrix, classes):
    workbook = Workbook()
    classification = workbook.active
    classification.title = "Classification_Report"
    classification.append(["Accuracy", report["accuracy"]])
    classification.append(["Samples", int(matrix.sum())])
    classification.append([])
    classification.append(["Label", "Precision", "Recall", "F1-score", "Support"])
    for label, metrics in report.items():
        if isinstance(metrics, dict):
            classification.append(
                [
                    label,
                    metrics["precision"],
                    metrics["recall"],
                    metrics["f1-score"],
                    metrics["support"],
                ]
            )
    classification["B1"].number_format = "0.00%"
    classification.freeze_panes = "B5"
    for row in classification.iter_rows(min_row=5, min_col=2, max_col=4):
        for cell in row:
            cell.number_format = "0.0000"

    confusion = workbook.create_sheet("Confusion_Matrix")
    confusion.append(["Actual / Predicted", *classes])
    for label, row in zip(classes, matrix):
        confusion.append([label, *row.tolist()])
    confusion.freeze_panes = "B2"
    for sheet, header_row in ((classification, 4), (confusion, 1)):
        for cell in sheet[header_row]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="244B5A")
        sheet.column_dimensions["A"].width = 28
        for column in range(2, sheet.max_column + 1):
            sheet.column_dimensions[get_column_letter(column)].width = (
                26 if sheet is confusion else 16
            )
    workbook.save(path)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--checkpoint", type=Path, required=True)
    parser.add_argument(
        "--data", type=Path, required=True, help="Held-out class folders"
    )
    parser.add_argument("--architecture", choices=ARCHITECTURES, required=True)
    parser.add_argument("--classes", type=Path, default=DEFAULT_CLASSES)
    parser.add_argument("--output", type=Path, default=Path("results/evaluation"))
    parser.add_argument("--image-size", type=int, default=224)
    parser.add_argument("--resize-size", type=int)
    parser.add_argument("--batch-size", type=int, default=16)
    parser.add_argument("--device", choices=("cpu", "cuda"), default="cpu")
    args = parser.parse_args()
    try:
        classes = read_classes(args.classes)
        transform = evaluation_transform(
            args.architecture, args.image_size, args.resize_size
        )
        dataset = ImageFolder(args.data, transform=transform)
        if dataset.classes != classes:
            raise ValueError(
                "Dataset class folders must match the class file in sorted order."
            )
        if args.batch_size <= 0:
            raise ValueError("Batch size must be positive.")
        if args.output.exists():
            raise ValueError("Output already exists; choose a new directory.")
        model = load_model(args.checkpoint, args.architecture, classes, args.device)
    except (ValueError, OSError) as error:
        parser.error(str(error))

    true_labels, predictions = [], []
    with torch.inference_mode():
        for images, labels in tqdm(
            DataLoader(dataset, batch_size=args.batch_size), desc="Evaluating"
        ):
            outputs = model(images.to(args.device))
            predictions.extend(outputs.argmax(dim=1).cpu().tolist())
            true_labels.extend(labels.tolist())
    indices = list(range(len(classes)))
    report = classification_report(
        true_labels,
        predictions,
        labels=indices,
        target_names=classes,
        output_dict=True,
        zero_division=0,
    )
    matrix = confusion_matrix(true_labels, predictions, labels=indices)
    args.output.mkdir(parents=True)
    metadata = {
        "architecture": args.architecture,
        "image_size": args.image_size,
        "resize_size": transform.transforms[0].size,
        "samples": len(dataset),
        "classes": classes,
        "metrics": report,
    }
    (args.output / "report.json").write_text(
        json.dumps(metadata, indent=2), encoding="utf-8"
    )
    with (args.output / "confusion_matrix.csv").open(
        "w", newline="", encoding="utf-8"
    ) as stream:
        writer = csv.writer(stream)
        writer.writerow(["actual/predicted", *classes])
        writer.writerows([name, *row.tolist()] for name, row in zip(classes, matrix))
    write_excel_report(args.output / "evaluation.xlsx", report, matrix, classes)
    print(
        classification_report(
            true_labels,
            predictions,
            labels=indices,
            target_names=classes,
            zero_division=0,
        )
    )
    print(f"Reports saved to {args.output}")


if __name__ == "__main__":
    main()
