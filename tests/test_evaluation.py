import tempfile
import unittest
from pathlib import Path

import numpy as np
from openpyxl import load_workbook
from sklearn.metrics import classification_report, confusion_matrix

from tools.evaluate import write_excel_report


class ExcelReportTests(unittest.TestCase):
    def test_workbook_preserves_confusion_counts_label_order_and_metrics(self):
        classes = ["other", "wood", "missing_prediction"]
        actual = [0, 1, 1, 2]
        predicted = [0, 0, 1, 1]
        indices = list(range(len(classes)))
        matrix = confusion_matrix(actual, predicted, labels=indices)
        report = classification_report(
            actual,
            predicted,
            labels=indices,
            target_names=classes,
            output_dict=True,
            zero_division=0,
        )
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "evaluation.xlsx"
            write_excel_report(output, report, matrix, classes)
            workbook = load_workbook(output, data_only=True)
            try:
                self.assertEqual(
                    workbook.sheetnames, ["Classification_Report", "Confusion_Matrix"]
                )
                sheet = workbook["Confusion_Matrix"]
                self.assertEqual(
                    [sheet.cell(1, col).value for col in range(2, 5)], classes
                )
                self.assertEqual(
                    [sheet.cell(row, 1).value for row in range(2, 5)], classes
                )
                values = [
                    [sheet.cell(row, col).value for col in range(2, 5)]
                    for row in range(2, 5)
                ]
                np.testing.assert_array_equal(values, matrix)
                self.assertEqual(sum(sum(row) for row in values), len(actual))
                classification = workbook["Classification_Report"]
                self.assertEqual(classification["B1"].value, 0.5)
                self.assertEqual(classification["B2"].value, len(actual))
                for row, label in enumerate(classes, start=5):
                    self.assertEqual(classification.cell(row, 1).value, label)
                    self.assertAlmostEqual(
                        classification.cell(row, 4).value, report[label]["f1-score"]
                    )
            finally:
                workbook.close()


if __name__ == "__main__":
    unittest.main()
